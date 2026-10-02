#!/usr/bin/env python3
"""Helpers for the jobs profile: ATS-safe resume docx + Telegram document send.

Usage (run by the jobs profile agent):
  python3 jobbot_helpers.py make-docx resume.md out.docx
  python3 jobbot_helpers.py send-doc out.docx "caption"
  python3 jobbot_helpers.py send-text "message"

resume.md format (markdown): # Name, ## sections, - bullets, plain lines.
Any tables/images/columns are stripped by construction — output is single-column ATS-safe.
"""
import sys
from pathlib import Path

PROJECT = Path.home() / "Hermes-workspace/projects/job-applications"
ENV = Path.home() / ".hermes/profiles/jobs/.env"


def load_token() -> str:
    for line in ENV.read_text().splitlines():
        if line.startswith("TELEGRAM_BOT_TOKEN="):
            return line.split("=", 1)[1].strip()
    raise SystemExit("TELEGRAM_BOT_TOKEN not found in jobs profile .env")


def load_chat_id() -> str:
    p = PROJECT / "cache" / "telegram_chat_id.txt"
    if not p.exists():
        raise SystemExit("chat_id unknown — user must /start the bot first; gateway records pairing")
    return p.read_text().strip()


def make_docx(src: str, dst: str) -> str:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def add(text: str, size=11, bold=False, center=False):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        return p

    for raw in Path(src).read_text().splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# ") and not line.startswith("## "):
            add(line[2:].strip(), size=16, bold=True, center=True)
        elif line.startswith("## "):
            add(line[3:].strip(), size=13, bold=True)
        elif line.lstrip().startswith(("- ", "• ")):
            doc.add_paragraph(line.lstrip()[2:].strip(), style="List Bullet")
        else:
            add(line.strip())
    out = Path(dst)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return str(out)


def tg(method: str, payload: dict) -> dict:
    import json, urllib.request
    url = f"https://api.telegram.org/bot{load_token()}/{method}"
    req = urllib.request.Request(url, data=json.dumps(payload).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read())


def send_text(text: str) -> dict:
    return tg("sendMessage", {"chat_id": load_chat_id(), "text": text[:4000]})


def send_doc(path: str, caption: str = "") -> dict:
    import uuid
    boundary = uuid.uuid4().hex
    body = b""
    fields = {"chat_id": load_chat_id()}
    if caption:
        fields["caption"] = caption[:1000]
    for k, v in fields.items():
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n").encode()
    fp = Path(path)
    body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"document\"; filename=\"{fp.name}\"\r\n"
             "Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n").encode()
    body += fp.read_bytes() + f"\r\n--{boundary}--\r\n".encode()
    import urllib.request
    req = urllib.request.Request(
        f"https://api.telegram.org/bot{load_token()}/sendDocument", data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    with urllib.request.urlopen(req, timeout=60) as r:
        import json
        return json.loads(r.read())


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "make-docx":
        print(make_docx(sys.argv[2], sys.argv[3]))
    elif cmd == "send-doc":
        print(send_doc(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else ""))
    elif cmd == "send-text":
        print(send_text(sys.argv[2]))
    else:
        raise SystemExit(__doc__)